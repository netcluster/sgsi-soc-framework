# -*- coding: utf-8 -*-
"""
Generador de la Guía Completa en Formato Microsoft Word (.docx) para el SGSI & SOC Framework.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo a una celda de tabla en Word"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_full_guide_docx(output_path: str):
    doc = Document()

    # Definir márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos de fuentes base
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # =========================================================================
    # PORTADA / TÍTULO PRINCIPAL
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("GUÍA DE IMPLEMENTACIÓN Y OPERACIÓN\nSGSI (ISO/IEC 27001:2022) & SOC")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Azul marino institucional

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Solución Integral Basada en Herramientas Gratuitas de Google y Motor Python\n")
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Meta-información en recuadro
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Estándar de Seguridad:", "ISO/IEC 27001:2022 & ISO/IEC 27005"),
        ("Tecnologías Utilizadas:", "Google Sheets, Looker Studio, Google Forms, Drive & Python 3"),
        ("Costo de Licenciamiento:", "100% Gratuito (Herramientas nativas Google Free Tier)"),
        ("Versión del Documento:", "1.0 - Edición Oficial")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1 = row.cells[0]
        c2 = row.cells[1]
        c1.text = k
        c1.paragraphs[0].runs[0].font.bold = True
        c2.text = v
        set_cell_background(c1, "F0F4F8")
        set_cell_background(c2, "FFFFFF")

    doc.add_paragraph("\n" + "="*55 + "\n")

    # =========================================================================
    # CAPÍTULO 1: INTRODUCCIÓN Y ARQUITECTURA
    # =========================================================================
    h1 = doc.add_heading("1. Introducción y Arquitectura de la Solución", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "El presente manual describe la puesta en marcha de un Sistema de Gestión de Seguridad de la Información (SGSI) "
        "completamente funcional y alineado a la norma internacional ISO/IEC 27001:2022, complementado con un Centro de "
        "Operaciones de Seguridad (SOC) para la detección, correlación y visualización de incidentes en tiempo real. "
        "Toda la arquitectura opera sin costo de licenciamiento utilizando el ecosistema de Google y un motor de automatización en Python."
    )

    doc.add_heading("Componentes de la Arquitectura:", level=2)
    p_comp = doc.add_paragraph()
    p_comp.add_run("1. Google Sheets (Base de Datos Central): ").bold = True
    p_comp.add_run("Alberga las 4 matrices del SGSI (Activos, Riesgos, SoA e Incidentes).\n")
    p_comp.add_run("2. Google Forms (Recolección): ").bold = True
    p_comp.add_run("Permite el reporte descentralizado de incidentes y solicitudes por parte de usuarios.\n")
    p_comp.add_run("3. Google Looker Studio (Dashboards): ").bold = True
    p_comp.add_run("Paneles de control ejecutivos y operativos en tiempo real.\n")
    p_comp.add_run("4. Motor Python (SOC & Analítica): ").bold = True
    p_comp.add_run("Procesa logs, detecta ataques, calcula riesgos según ISO 27005 y sincroniza los datos.\n")

    # =========================================================================
    # CAPÍTULO 2: REQUISITOS E INSTALACIÓN
    # =========================================================================
    h2 = doc.add_heading("2. Requisitos Previos e Instalación", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "Para ejecutar el motor del framework, únicamente se requiere un entorno estándar con Python 3.9 o superior y "
        "una cuenta estándar de Google (Gmail o Google Workspace)."
    )

    doc.add_heading("Estructura de Carpetas del Proyecto:", level=2)
    struct_p = doc.add_paragraph()
    struct_p.add_run(
        "sgsi_soc_framework/\n"
        "├── config/                 -> Configuración, umbrales y catálogo de 93 controles ISO 27001\n"
        "├── core/                   -> Motores de cálculo de riesgos, incidentes y sincronización Google\n"
        "├── soc_engine/             -> Parsers de logs, detector de amenazas y mapeador MITRE ATT&CK\n"
        "├── templates_google/       -> Matrices base en formato CSV (UTF-8 BOM)\n"
        "├── sync_drive/             -> Archivos listos para subir a Google Drive\n"
        "├── simulators/             -> Generadores de ataques y telemetría de prueba\n"
        "├── dashboards/             -> Guías y scripts para Google Looker Studio y Apps Script\n"
        "└── main.py                 -> CLI principal de gestión\n"
    )

    doc.add_heading("Instalación de Dependencias (Opcional):", level=2)
    doc.add_paragraph(
        "El núcleo funciona 100% con las bibliotecas estándar de Python. Si deseas usar la API directa de Google Cloud, instala:"
    )
    cmd_p = doc.add_paragraph()
    cmd_p.add_run("pip install gspread oauth2client pyyaml\n").bold = True

    # =========================================================================
    # CAPÍTULO 3: MÉTODOS DE SINCRONIZACIÓN CON GOOGLE
    # =========================================================================
    h3 = doc.add_heading("3. Métodos de Sincronización con Google (Dual-Mode)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "El framework cuenta con dos métodos de conexión diseñados para adaptarse a la disponibilidad de credenciales de Google Cloud Platform (GCP):"
    )

    # Tabla comparativa
    table_sync = doc.add_table(rows=3, cols=3)
    table_sync.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_sync = ["Característica", "Método A: Google Cloud API", "Método B: Local Sync (Sin GCP)"]
    for j, h in enumerate(headers_sync):
        cell = table_sync.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "1B365D")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_sync = [
        ("Requisitos", "Service Account JSON de Google Cloud (Free Tier)", "Ninguno (Solo cuenta de Google/Drive)"),
        ("Automatización", "100% Nube automática mediante API de Google Sheets", "Exportación CSV + Google Apps Script en 1 clic")
    ]
    for row_idx, row_data in enumerate(data_sync):
        for col_idx, text in enumerate(row_data):
            cell = table_sync.rows[row_idx + 1].cells[col_idx]
            cell.text = text
            set_cell_background(cell, "F9FBFD" if row_idx % 2 == 0 else "FFFFFF")

    doc.add_paragraph("\n")
    doc.add_heading("Paso a Paso - Método B: Local Sync (Recomendado sin GCP):", level=2)
    p_steps_b = doc.add_paragraph()
    p_steps_b.add_run("1. Ejecuta el comando de sincronización:\n").bold = True
    p_steps_b.add_run("   python main.py --action sync\n")
    p_steps_b.add_run("2. Los archivos actualizados se generarán en la carpeta ").bold = True
    p_steps_b.add_run("sync_drive/\n")
    p_steps_b.add_run("3. Sube los archivos a tu Google Drive o utiliza la herramienta Google Drive Desktop para sincronización transparente.\n")
    p_steps_b.add_run("4. En tu Google Sheets, abre Extensiones > Apps Script, pega el archivo ").bold = True
    p_steps_b.add_run("dashboards/google_apps_script_sync.js y haz clic en Ejecutar para poblar todas las hojas en 1 clic.\n")

    # =========================================================================
    # CAPÍTULO 4: OPERACIÓN DEL SGSI (ISO/IEC 27001 & ISO 27005)
    # =========================================================================
    h4 = doc.add_heading("4. Operación del SGSI (Gestión de Seguridad de la Información)", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "El SGSI se gestiona a través de 3 pilares documentales estructurados en Google Sheets y procesados por el motor de cálculo en Python:"
    )

    p_pilares = doc.add_paragraph()
    p_pilares.add_run("A. Inventario y Valoración de Activos:\n").bold = True
    p_pilares.add_run("   - Calificación de Confidencialidad, Integridad y Disponibilidad en escala de 1 a 5.\n"
                      "   - Cálculo automático del nivel de criticidad del activo (Crítico, Alto, Medio, Bajo).\n\n")

    p_pilares.add_run("B. Matriz de Riesgos (ISO 27005):\n").bold = True
    p_pilares.add_run("   - Evaluación de Riesgo Inherente = Probabilidad × Impacto.\n"
                      "   - Evaluación de Eficacia de Controles (% de mitigación).\n"
                      "   - Cálculo de Riesgo Residual y determinación de la estrategia (Mitigar, Aceptar, Transferir, Evitar).\n"
                      "   - Comando de recálculo: ").bold = False
    p_pilares.add_run("python main.py --action calculate-risks\n\n")

    p_pilares.add_run("C. Declaración de Aplicabilidad (SoA - ISO 27001:2022):\n").bold = True
    p_pilares.add_run("   - Seguimiento integral a los 93 controles distribuidos en 4 dominios:\n"
                      "     1. Controles Organizacionales (37 controles)\n"
                      "     2. Controles de Personas (8 controles)\n"
                      "     3. Controles Físicos (14 controles)\n"
                      "     4. Controles Tecnológicos (34 controles)\n")

    # =========================================================================
    # CAPÍTULO 5: OPERACIÓN DEL SOC Y TELEMETRÍA
    # =========================================================================
    h5 = doc.add_heading("5. Operación del SOC (Centro de Operaciones de Seguridad)", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "El motor del SOC ingesta eventos de seguridad, ejecuta reglas de detección de amenazas y mapea automáticamente "
        "cada incidente con las técnicas y tácticas del marco internacional MITRE ATT&CK."
    )

    doc.add_heading("Reglas de Detección Incorporadas:", level=2)
    p_rules = doc.add_paragraph()
    p_rules.add_run("• Ataques de Fuerza Bruta (SSH / Auth): ").bold = True
    p_rules.add_run("Detecta intentos fallidos reiterados en ventanas de tiempo cortas (MITRE T1110).\n")
    p_rules.add_run("• Inyecciones SQL (SQLi): ").bold = True
    p_rules.add_run("Identifica patrones en URLs y peticiones web (UNION, SELECT, OR 1=1) (MITRE T1190).\n")
    p_rules.add_run("• Cross-Site Scripting (XSS): ").bold = True
    p_rules.add_run("Detecta scripts y etiquetas HTML sospechosas en peticiones web (MITRE T1189).\n")
    p_rules.add_run("• Evasión y Malware en Endpoints: ").bold = True
    p_rules.add_run("Detecta ejecuciones de PowerShell ofuscadas en Base64 o con políticas de bypass (MITRE T1059).\n")

    doc.add_heading("Comandos del SOC:", level=2)
    p_soc_cmds = doc.add_paragraph()
    p_soc_cmds.add_run("• Simular flujo de ataques para pruebas: ").bold = True
    p_soc_cmds.add_run("python main.py --action simulate-soc\n")
    p_soc_cmds.add_run("• Ingestar un archivo de logs real: ").bold = True
    p_soc_cmds.add_run("python main.py --action ingest-logs --log-file ruta/a/logs.log\n")
    p_soc_cmds.add_run("• Ver KPIs de rendimiento (MTTD, MTTR): ").bold = True
    p_soc_cmds.add_run("python main.py --action show-kpis\n")

    # =========================================================================
    # CAPÍTULO 6: CONSTRUCCIÓN DE DASHBOARDS EN LOOKER STUDIO
    # =========================================================================
    h6 = doc.add_heading("6. Construcción de Dashboards en Google Looker Studio", level=1)
    h6.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "Google Looker Studio (100% gratuito) permite crear paneles visuales que se actualizan de manera continua "
        "al consultar directamente las hojas de Google Sheets."
    )

    doc.add_heading("Paso 1: Conectar Google Sheets a Looker Studio", level=2)
    doc.add_paragraph(
        "1. Ingresa a https://lookerstudio.google.com/ con tu cuenta de Google.\n"
        "2. Haz clic en Crear > Informe y selecciona el conector Google Sheets.\n"
        "3. Selecciona tu hoja maestra e incorpora las 4 pestañas como fuentes de datos."
    )

    doc.add_heading("Paso 2: Configuración del Dashboard Operativo SOC", level=2)
    doc.add_paragraph(
        "• Tarjeta MTTD Promedio: Métrica Promedio(MTTD_Minutos).\n"
        "• Tarjeta MTTR Promedio: Métrica Promedio(MTTR_Minutos) con filtro Estado = 'Cerrado'.\n"
        "• Gráfico Circular de Severidades: Dimensión Severidad, Métrica Recuento(ID_Incidente).\n"
        "• Gráfico de Barras de Amenazas: Dimensión Tipo_Amenaza, Métrica Recuento(ID_Incidente).\n"
        "• Matriz MITRE ATT&CK: Eje Y Tactica_MITRE, desglose Tecnica_MITRE.\n"
        "• Tabla de Alertas en Vivo: Columnas Fecha_Hora, Titulo_Incidente, Severidad, Estado, IP_Origen, Accion_Correctiva."
    )

    doc.add_heading("Paso 3: Configuración del Dashboard Ejecutivo SGSI", level=2)
    doc.add_paragraph(
        "• Indicador de Madurez ISO 27001: Métrica Promedio(Porcentaje_Madurez) de la hoja SoA.\n"
        "• Cumplimiento por Dominio: Gráfico de barras apiladas con Dimensión Dominio y desglose Estado_Implementacion.\n"
        "• Mapa de Calor de Riesgos ISO 27005: Gráfico de dispersión Probabilidad_Residual_1a5 vs Impacto_Residual_1a5."
    )

    # =========================================================================
    # CAPÍTULO 7: AUTOMATIZACIÓN Y PROGRAMACIÓN PERIÓDICA
    # =========================================================================
    h7 = doc.add_heading("7. Automatización Continua y Mantenimiento", level=1)
    h7.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "Para mantener el SGSI y SOC operando de forma desatendida, se recomienda programar la ejecución periódica del comando unificado:"
    )

    p_cron = doc.add_paragraph()
    p_cron.add_run("python main.py --action all\n").bold = True

    doc.add_heading("Configuración en Programador de Tareas de Windows:", level=2)
    doc.add_paragraph(
        "1. Abre el Programador de Tareas de Windows (taskschd.msc).\n"
        "2. Crea una Tarea Básica llamada 'SGSI_SOC_Sync'.\n"
        "3. Desencadenador: Diariamente o Cada 1 hora.\n"
        "4. Acción: Iniciar un programa.\n"
        "   - Programa o script: python.exe\n"
        "   - Argumentos: main.py --action all\n"
        "   - Iniciar en: C:\\Users\\RICARDO.ALFARO\\.gemini\\antigravity\\scratch\\sgsi_soc_framework\n"
    )

    # =========================================================================
    # CAPÍTULO 8: MATRIZ DE COMANDOS RÁPIDOS
    # =========================================================================
    h8 = doc.add_heading("8. Resumen de Comandos Rápidos del CLI", level=1)
    h8.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    table_cmd = doc.add_table(rows=6, cols=2)
    table_cmd.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_headers = ["Acción Deseada", "Comando a Ejecutar"]
    for j, h in enumerate(cmd_headers):
        cell = table_cmd.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "1B365D")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    cmd_rows = [
        ("Ejecutar todo el ciclo (Riesgos + Sincronización + KPIs)", "python main.py --action all"),
        ("Sincronizar matrices para Google Drive", "python main.py --action sync"),
        ("Sincronizar vía API directa de Google Cloud", "python main.py --action sync --sync-mode service_account --share-email tu_correo@gmail.com"),
        ("Recalcular Matriz de Riesgos ISO 27005", "python main.py --action calculate-risks"),
        ("Simular ataques y telemetría SOC en vivo", "python main.py --action simulate-soc"),
    ]
    for row_idx, (act, cmd) in enumerate(cmd_rows):
        row = table_cmd.rows[row_idx + 1]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = act
        c2.text = cmd
        c2.paragraphs[0].runs[0].font.name = 'Consolas'
        c2.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c1, "F9FBFD" if row_idx % 2 == 0 else "FFFFFF")
        set_cell_background(c2, "F9FBFD" if row_idx % 2 == 0 else "FFFFFF")

    doc.save(output_path)
    print(f"[Word Guide Generator] Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    out_file = r"C:\Users\RICARDO.ALFARO\.gemini\antigravity\scratch\sgsi_soc_framework\GUIA_IMPLEMENTACION_SGSI_SOC_WORD.docx"
    create_full_guide_docx(out_file)
