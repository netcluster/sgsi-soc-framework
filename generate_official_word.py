# -*- coding: utf-8 -*-
"""
Script generador del documento nativo oficial de Microsoft Word (.docx).
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_official_word_doc(output_paths):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(44, 62, 80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("GUÍA DE IMPLEMENTACIÓN Y OPERACIÓN\nSGSI (ISO/IEC 27001:2022) & SOC")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Solución Integral con Herramientas Gratuitas de Google y Motor Python\n")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(120, 120, 120)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Norma de Referencia:", "ISO/IEC 27001:2022 & ISO/IEC 27005"),
        ("Ecosistema Utilizado:", "Google Sheets, Google Looker Studio, Google Forms, Drive & Python"),
        ("Ubicación del Proyecto:", r"C:\Users\RICARDO.ALFARO\Documents\sgsi_soc_framework"),
        ("Estado de Implementación:", "Completamente Funcional y Listo para Despliegue")
    ]
    for idx, (label, val) in enumerate(meta_info):
        row = table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.text = label
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        cell_val.text = val
        shd_lbl = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
        shd_val = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell_lbl._tc.get_or_add_tcPr().append(shd_lbl)
        cell_val._tc.get_or_add_tcPr().append(shd_val)

    doc.add_paragraph("\n")

    h1 = doc.add_heading("1. Introducción y Arquitectura de la Solución", level=1)
    h1.runs[0].font.color.rgb = RGBColor(27, 54, 93)
    
    doc.add_paragraph(
        "Este manual describe paso a paso la implementación y uso de un Sistema de Gestión de Seguridad "
        "de la Información (SGSI) alineado a ISO/IEC 27001:2022 y un Centro de Operaciones de Seguridad (SOC) "
        "para el monitoreo y respuesta a ciberataques en tiempo real."
    )
    
    p_comp = doc.add_paragraph()
    p_comp.add_run("• Google Sheets: ").bold = True
    p_comp.add_run("Base de datos relacional con 4 matrices (Activos, Riesgos, SoA e Incidentes).\n")
    p_comp.add_run("• Google Forms: ").bold = True
    p_comp.add_run("Formularios para reporte descentralizado de incidentes y alta de activos.\n")
    p_comp.add_run("• Google Looker Studio: ").bold = True
    p_comp.add_run("Dashboards ejecutivos y operativos en tiempo real.\n")
    p_comp.add_run("• Motor Python: ").bold = True
    p_comp.add_run("Ingesta de logs, correlación con MITRE ATT&CK, cálculo de riesgos ISO 27005 y sincronización.\n")

    h2 = doc.add_heading("2. Métodos de Sincronización con Google (Dual-Mode)", level=1)
    h2.runs[0].font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph(
        "El framework soporta dos formas de trabajo según si tienes o no credenciales de Google Cloud Platform (GCP):"
    )

    p_sync = doc.add_paragraph()
    p_sync.add_run("A. Modo Directo (Google Cloud API):\n").bold = True
    p_sync.add_run(
        "Si dispones de cuenta de Google Cloud (Free Tier), descarga tu Service Account JSON en config/google_credentials.json. "
        "El sistema creará la hoja automáticamente y actualizará las celdas en la nube.\n"
        "Comando: python main.py --action sync --sync-mode service_account --share-email tu_correo@gmail.com\n\n"
    )
    p_sync.add_run("B. Modo Local Sync (100% Sin Google Cloud):\n").bold = True
    p_sync.add_run(
        "Si no tienes cuenta de GCP, ejecuta: python main.py --action sync\n"
        "Se generarán los 4 archivos CSV en la carpeta 'sync_drive/'. Puedes subirlos a Google Drive e importarlos "
        "en 1 solo clic pegando el script 'dashboards/google_apps_script_sync.js' en Extensiones > Apps Script de Google Sheets.\n"
    )

    h3 = doc.add_heading("3. Operación del SGSI (ISO 27001 & ISO 27005)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph(
        "1. Inventario de Activos: Clasificación de Confidencialidad, Integridad y Disponibilidad (1 a 5) y cálculo de criticidad.\n"
        "2. Matriz de Riesgos (ISO 27005): Evaluación de Riesgo Inherente (Probabilidad x Impacto) y reducción a Riesgo Residual según eficacia de controles.\n"
        "3. Declaración de Aplicabilidad (SoA): Monitoreo de madurez de los 93 controles en sus 4 dominios (Organizacional, Personas, Físico, Tecnológico).\n"
        "Comando para recalcular riesgos: python main.py --action calculate-risks"
    )

    h4 = doc.add_heading("4. Operación del SOC y Detección de Amenazas", level=1)
    h4.runs[0].font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph(
        "El motor del SOC analiza logs de autenticación, servidores web y endpoints, correlacionando eventos contra MITRE ATT&CK:\n"
        "• Fuerza Bruta (SSH / Login) -> MITRE T1110\n"
        "• Inyección SQL (SQLi) -> MITRE T1190\n"
        "• Cross-Site Scripting (XSS) -> MITRE T1189\n"
        "• Evasión / Scripts PowerShell -> MITRE T1059\n\n"
        "Comandos del SOC:\n"
        "• Simular ciberataques en vivo: python main.py --action simulate-soc\n"
        "• Ingestar un log real: python main.py --action ingest-logs --log-file ruta/a/logs.log\n"
        "• Ver métricas operativas (MTTD / MTTR): python main.py --action show-kpis"
    )

    h5 = doc.add_heading("5. Despliegue de Dashboards en Vivo y Looker Studio", level=1)
    h5.runs[0].font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph(
        "1. Dashboards Interactivos en Python (Tiempo Real): Ejecuta 'python dashboards/live_server.py' y abre http://localhost:8080/ciso o http://localhost:8080/direccion.\n"
        "2. Looker Studio: Conecta tu hoja de Google Sheets en https://lookerstudio.google.com/ para compartir reportes en la nube."
    )

    h6 = doc.add_heading("6. Resumen de Comandos del CLI", level=1)
    h6.runs[0].font.color.rgb = RGBColor(27, 54, 93)

    cmd_table = doc.add_table(rows=7, cols=2)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Acción", "Comando"]
    for j, h in enumerate(headers):
        c = cmd_table.rows[0].cells[j]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
        c._tc.get_or_add_tcPr().append(shd)

    rows_data = [
        ("Ejecutar ciclo completo", "python main.py --action all"),
        ("Servidor de Dashboards en Vivo", "python dashboards/live_server.py"),
        ("Sincronizar datos con Google", "python main.py --action sync"),
        ("Recalcular Matriz de Riesgos", "python main.py --action calculate-risks"),
        ("Simular ataques SOC en vivo", "python main.py --action simulate-soc"),
        ("Mostrar KPIs en consola", "python main.py --action show-kpis")
    ]
    for r_i, (act, cmd) in enumerate(rows_data):
        row = cmd_table.rows[r_i + 1]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = act
        c2.text = cmd
        c2.paragraphs[0].runs[0].font.name = 'Consolas'
        c2.paragraphs[0].runs[0].font.size = Pt(9.5)
        bg = "F9FBFD" if r_i % 2 == 0 else "FFFFFF"
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
        shd2 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
        c1._tc.get_or_add_tcPr().append(shd1)
        c2._tc.get_or_add_tcPr().append(shd2)

    for out in output_paths:
        doc.save(out)
        print(f"[Word Docx Creator] Documento Word guardado en: {out}")

if __name__ == "__main__":
    paths = [
        r"C:\Users\RICARDO.ALFARO\Desktop\GUIA_SGSI_SOC_ISO27001.docx",
        r"C:\Users\RICARDO.ALFARO\Documents\sgsi_soc_framework\GUIA_SGSI_SOC_ISO27001.docx"
    ]
    build_official_word_doc(paths)
